#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""一次性脚本：提取已填写的 Project Verification Form 的段落 + 表格 + 复选框状态。
输出到 stdout（UTF-8）。仅读取，不修改原文件。
"""
import sys, io, os, zipfile, re

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

SRC = r"C:\Users\<you>\Downloads\Project Verification Form_NiYoujia (for own project)_Final.docx"

if not os.path.exists(SRC):
    print("!! FILE NOT FOUND:", SRC)
    sys.exit(1)

print("=" * 78)
print("FILE:", SRC)
print("SIZE:", os.path.getsize(SRC))
print("=" * 78)

try:
    import docx
    HAVE_DOCX = True
except Exception as e:
    HAVE_DOCX = False
    print("!! python-docx unavailable:", e)

if HAVE_DOCX:
    d = docx.Document(SRC)

    print("\n----- BODY ELEMENTS IN ORDER (paragraphs + tables) -----\n")
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    body = d.element.body
    pi = ti = 0
    for child in body.iterchildren():
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            p = Paragraph(child, d)
            txt = p.text
            if txt.strip():
                print("P%03d | %s" % (pi, txt))
            else:
                print("P%03d | <empty>" % pi)
            pi += 1
        elif tag == 'tbl':
            t = Table(child, d)
            print("\n### TABLE %d  (rows=%d)" % (ti, len(t.rows)))
            for ri, row in enumerate(t.rows):
                cells = []
                for c in row.cells:
                    cells.append(c.text.replace("\n", " / ").strip())
                print("  T%d R%d | %s" % (ti, ri, " || ".join(cells)))
            print()
            ti += 1

    print("\n----- CHECKBOX / TICK CHARACTER SCAN (all text) -----")
    alltext = []
    for p in d.paragraphs:
        alltext.append(p.text)
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                alltext.append(c.text)
    blob = "\n".join(alltext)
    for ch, name in [("\u2610", "BALLOT BOX"),
                     ("\u2611", "BALLOT BOX WITH CHECK"),
                     ("\u2612", "BALLOT BOX WITH X"),
                     ("\u25a1", "WHITE SQUARE"),
                     ("\u25aa", "BLACK SMALL SQUARE"),
                     ("\u2713", "CHECK MARK"),
                     ("\u2714", "HEAVY CHECK MARK"),
                     ("\u00d7", "MULTIPLICATION SIGN"),
                     ("\u221a", "SQRT/TICK")]:
        n = blob.count(ch)
        if n:
            print("  %s U+%04X %s x%d" % (ch, ord(ch), name, n))
    print("  (chars not listed above = count 0)")

    print("\n----- SECTION 1..7 BLANK FIELD CHECK -----")
    for p in d.paragraphs:
        s = p.text.strip()
        if re.match(r'^[A-Z][A-Za-z /()\']+:\s*$', s):
            print("  BLANK-LOOKING FIELD: %r" % s)

print("\n----- RAW document.xml: content controls / checkbox markers -----")
with zipfile.ZipFile(SRC) as z:
    names = z.namelist()
    print("  parts:", [n for n in names if n.endswith('.xml')][:20])
    xml = z.read('word/document.xml').decode('utf-8', 'replace')
    print("  len(document.xml) =", len(xml))
    for marker in ['w:checkBox', 'w:sdt', 'w14:checkbox', 'FORMCHECKBOX', 'w:sym',
                   'w:ffData', 'w:default', 'w:checked']:
        print("  %-16s x%d" % (marker, xml.count(marker)))
print("\nDONE")
