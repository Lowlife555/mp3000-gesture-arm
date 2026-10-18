#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""一次性脚本：解析已填表里 w14:checkbox 的勾选状态 + 逐字计数 + 原文 repr 核对。只读。"""
import sys, io, os, zipfile, re, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

SRC = r"C:\Users\<you>\Downloads\Project Verification Form_NiYoujia (for own project)_Final.docx"
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
W   = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

import xml.etree.ElementTree as ET
import docx

d = docx.Document(SRC)

print("=" * 78)
print("A. 每个含 checkbox 的段落 —— 勾选状态")
print("=" * 78)
print("规则: <w14:checked w14:val=\"1\"/> = 已勾；val=\"0\" 或缺 = 未勾\n")

for i, p in enumerate(d.paragraphs):
    xml = p._p.xml
    if 'checkbox' not in xml:
        continue
    # 段落文字
    txt = p.text.strip()
    # 找出该段落里所有 checkbox 及其 checked 值
    states = []
    for m in re.finditer(r'<w14:checkbox>(.*?)</w14:checkbox>', xml, re.S):
        inner = m.group(1)
        cm = re.search(r'<w14:checked[^>]*w14:val="([^"]+)"', inner)
        if not cm:
            states.append("NO-VAL")
        else:
            states.append(cm.group(1))
    print("P%03d  checked=%s" % (i, states))
    print("       text=%r" % txt)

print()
print("=" * 78)
print("B. 全局统计")
print("=" * 78)
with zipfile.ZipFile(SRC) as z:
    xmlall = z.read('word/document.xml').decode('utf-8', 'replace')
tot = len(re.findall(r'<w14:checkbox>', xmlall))
on  = len(re.findall(r'<w14:checked[^>]*w14:val="(?:1|true|on)"', xmlall))
off = len(re.findall(r'<w14:checked[^>]*w14:val="(?:0|false|off)"', xmlall))
print("  checkbox 总数 =", tot)
print("  已勾(val=1) =", on)
print("  未勾(val=0) =", off)
print("  无 val      =", tot - on - off)

print()
print("=" * 78)
print("C. 字数核对（表单上限 150）")
print("=" * 78)
def wc(t):
    return len(re.findall(r"[A-Za-z0-9][A-Za-z0-9'\-/]*", t))

paras = [p.text for p in d.paragraphs]
def block(start_key, stop_key):
    """取 start_key 之后、stop_key 之前的全部非空段落（真正的栏位内容）。"""
    out, on = [], False
    for t in paras:
        s = t.strip()
        if s.startswith(start_key):
            on = True
            continue
        if on and s.startswith(stop_key):
            break
        if on and s:
            out.append(s)
    return out

sec2 = block("Brief Project Summary", "Project Origin")
sec5 = [t for t in paras if t.strip().startswith("This project sits within")]
print("  §2 Brief Project Summary 段落数 =", len(sec2), " 合计词数 =", sum(wc(t) for t in sec2))
for t in sec2:
    print("      (%d 词) %s" % (wc(t), t[:70] + "..."))
print()
print("  §5 Course Relevance 段落数 =", len(sec5), " 合计词数 =", sum(wc(t) for t in sec5))
for t in sec5:
    print("      (%d 词) %s" % (wc(t), t[:70] + "..."))

print()
print("=" * 78)
print("D. §4 表格：单元格原文 repr（确认 ' / ' 是不是我的提取产物）")
print("=" * 78)
for ti, t in enumerate(d.tables):
    for ri, row in enumerate(t.rows):
        for ci, c in enumerate(row.cells):
            print("  T%d R%d C%d = %r" % (ti, ri, ci, c.text))

print()
print("=" * 78)
print("E. §5 里点名的 programme 名称 vs Section 1 填的 programme")
print("=" * 78)
for t in paras:
    if 'Registered Course' in t or 'Programme:' in t:
        print("  SECTION1 行: %r" % t)
for t in paras:
    if 'Electrical, Electronics' in t:
        print("  学籍名出现: %r" % t)
    if 'electronic and electrical engineering' in t.lower():
        print("  §5 里的说法: %r" % t[:120])

print()
print("=" * 78)
print("F. 逐栏空值核对（哪些该填的没填）")
print("=" * 78)
checks = {
    "Student Name":        "Student Name:",
    "Student ID":          "Student ID:",
    "Registered Course":   "Registered Course / Programme:",
    "Mode of Study":       "Mode of Study:",
    "Module Code":         "Project Module Code:",
    "Academic Year":       "Academic Year:",
    "Project Title":       "Proposed Project Title:",
    "Supervisor Name":     "Proposed Supervisor Name:",
}
for label, key in checks.items():
    hit = [t for t in paras if t.strip().startswith(key)]
    if hit:
        inline = hit[0].strip()[len(key):].strip()
        idx = paras.index(hit[0])
        nxt = paras[idx+1].strip() if idx+1 < len(paras) else ""
        print("  %-18s inline=%r  next-line=%r" % (label, inline, nxt[:70]))
    else:
        print("  %-18s !! 标签未找到" % label)
print("\nDONE")
